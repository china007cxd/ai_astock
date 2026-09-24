# -*- coding: utf-8 -*-
"""CSV 与带样式 Excel 导出。"""
from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import settings
from .schemas import ExportRequest


def _safe_name(value: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "_", value.strip())[:60]
    return name or "A股数据导出"


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return str(value)
    return "" if value is None else value


def export_data(request: ExportRequest) -> Path:
    settings.export_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = settings.export_dir / f"{_safe_name(request.title)}_{stamp}.{request.format}"
    headers = [column.get("title") or column.get("key") or "" for column in request.columns]
    keys = [column.get("key") or "" for column in request.columns]
    if request.format == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(headers)
            writer.writerows([[_cell(row.get(key)) for key in keys] for row in request.rows])
        return path

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_name(request.title)[:31]
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1677FF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in request.rows:
        sheet.append([_cell(row.get(key)) for key in keys])
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        samples = [str(sheet.cell(row=row, column=index).value or "") for row in range(1, min(sheet.max_row, 101) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(10, max(map(len, samples), default=len(header)) + 2))
    workbook.save(path)
    return path
